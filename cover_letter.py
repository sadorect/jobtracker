"""Generate a draft cover letter (.docx) for a tracked role.

The app runs offline, so this produces a structured *draft* — pre-filled with the
role's company, title and your fit notes, with bracketed placeholders to finish.
For a fully tailored letter, hand the role to Claude / the job-hunt skill, which
has your profile.
"""

from __future__ import annotations

import datetime as dt
import re
from pathlib import Path

from docx import Document

import config


def _safe(name: str) -> str:
    """Make a string safe to use in a filename."""
    cleaned = re.sub(r'[\\/:*?"<>|]+', "", (name or "")).strip()
    return cleaned or "Untitled"


def cover_letter_path(job: dict, out_dir: Path | str | None = None) -> Path:
    out_dir = Path(out_dir) if out_dir else config.COVER_LETTER_DIR
    company = _safe(job.get("company") or "Company")
    title = _safe(job.get("title") or "Role")
    return out_dir / f"Cover Letter - {company} - {title}.docx"


def generate_cover_letter(job: dict, out_dir: Path | str | None = None) -> Path:
    """Write a draft cover letter for ``job`` and return the saved path."""
    out_dir = Path(out_dir) if out_dir else config.COVER_LETTER_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    path = cover_letter_path(job, out_dir)

    company = (job.get("company") or "the company").strip()
    title = (job.get("title") or "the role").strip()
    location = (job.get("location") or "").strip()
    fit = (job.get("fit_notes") or "").strip()
    today = dt.date.today().strftime("%d %B %Y")

    doc = Document()
    doc.add_paragraph(today)
    doc.add_paragraph("")
    doc.add_paragraph("Dear Hiring Manager,")

    intro = (
        f"I am writing to apply for the {title} position at {company}"
        + (f" ({location})" if location else "")
        + ". Having reviewed the role, I am confident that my background and "
        "experience make me a strong fit."
    )
    doc.add_paragraph(intro)

    body = (
        "In my career to date, [add two or three sentences on your most relevant "
        "experience, skills and achievements for this role]."
    )
    if fit:
        body += f" This role stood out to me in particular because {fit[0].lower()}{fit[1:]}."
    doc.add_paragraph(body)

    doc.add_paragraph(
        f"I would welcome the opportunity to discuss how I can contribute to {company}. "
        "Thank you for considering my application; I look forward to hearing from you."
    )
    doc.add_paragraph("")
    doc.add_paragraph("Yours sincerely,")
    doc.add_paragraph("[Your name]")

    doc.save(str(path))
    return path


def write_ai_cover_letter(job: dict, letter: dict, out_dir: Path | str | None = None) -> Path:
    """Save a cover letter Claude drafted (the claude-web ``cover_letter`` result).

    Uses the same name as the template draft plus a ``(Claude)`` suffix, so a
    template you have already edited is never overwritten.
    """
    out_dir = Path(out_dir) if out_dir else config.COVER_LETTER_DIR
    out_dir.mkdir(parents=True, exist_ok=True)
    template = cover_letter_path(job, out_dir)
    path = template.with_name(f"{template.stem} (Claude){template.suffix}")

    doc = Document()
    doc.add_paragraph(dt.date.today().strftime("%d %B %Y"))
    if letter.get("subject"):
        doc.add_paragraph(f"Re: {letter['subject']}")
    doc.add_paragraph("")
    doc.add_paragraph(letter.get("greeting") or "Dear Hiring Manager,")
    for para in letter.get("paragraphs") or []:
        doc.add_paragraph(para)
    doc.add_paragraph("")
    for line in (letter.get("sign_off") or "Yours sincerely,").splitlines():
        doc.add_paragraph(line)
    doc.save(str(path))
    return path


def letter_text(letter: dict) -> str:
    """The drafted letter as plain text, for previewing or pasting into an email."""
    parts = [letter.get("greeting") or "", *(letter.get("paragraphs") or []),
             letter.get("sign_off") or ""]
    return "\n\n".join(p for p in parts if p)
