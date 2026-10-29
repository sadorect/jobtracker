import io

import pytest
from docx import Document

import cv_upload

BODY = "Ada Example\nBusiness Analyst\n" + "Led requirements for a claims platform migration.\n" * 8


def _docx_bytes() -> bytes:
    doc = Document()
    for line in BODY.splitlines():
        doc.add_paragraph(line)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skills"
    table.cell(0, 1).text = "BPMN, SQL"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _pdf_bytes(text: str) -> bytes:
    """A minimal one-page text PDF (no extra dependencies)."""
    lines = text.splitlines()
    stream = "BT /F1 10 Tf 40 800 Td 12 TL " + " ".join(f"({ln}) '" for ln in lines) + " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out, offsets = "%PDF-1.4\n", []
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{i} 0 obj\n{obj}\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n"
    out += "".join(f"{o:010d} 00000 n \n" for o in offsets)
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    return out.encode("latin-1")


def test_docx_text_includes_tables():
    text = cv_upload.extract_text("cv.docx", _docx_bytes())
    assert text.startswith("Ada Example")
    assert "Skills | BPMN, SQL" in text


def test_pdf_text():
    text = cv_upload.extract_text("cv.pdf", _pdf_bytes(BODY))
    assert "Ada Example" in text and "claims platform migration" in text


@pytest.mark.parametrize(
    ("name", "data", "match"),
    [
        ("cv.exe", b"x" * 500, "Upload a"),
        ("cv.txt", b"too short", "Hardly any text"),
        ("cv.docx", b"not a zip" * 100, "Couldn't read that Word file"),
        ("cv.pdf", b"%PDF-garbage" * 50, "PDF"),
        ("cv.txt", b"x" * (cv_upload.MAX_BYTES + 1), "over 5 MB"),
    ],
)
def test_rejects_unusable_files(name, data, match):
    with pytest.raises(cv_upload.CvUploadError, match=match):
        cv_upload.extract_text(name, data)


def test_save_keeps_every_upload_and_tracks_the_latest(tmp_path):
    assert cv_upload.current_cv(tmp_path) is None
    first = cv_upload.save_upload("Old CV.docx", _docx_bytes(), tmp_path)
    assert (tmp_path / first["file"]).read_bytes() == _docx_bytes()
    second = cv_upload.save_upload("new.txt", BODY.replace("Ada", "Ada B").encode(), tmp_path)
    current = cv_upload.current_cv(tmp_path)
    assert current["original_name"] == "new.txt" and current["text"].startswith("Ada B Example")
    assert current["file"] == second["file"]
    assert (tmp_path / second["text_file"]).name.endswith(".extracted.txt")
