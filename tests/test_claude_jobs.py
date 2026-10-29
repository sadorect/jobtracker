"""claude_jobs: talking to a Claude Code Web jobs API and filing results back.

A small in-process HTTP server stands in for claude-web.
"""

import json
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest
from docx import Document

import claude_jobs as cj
import config
import cv_upload
import db

PROFILE = {
    "name": "Ada Example",
    "contact": "London  |  ada@example.com",
    "profiles": {"default": "Business analyst."},
    "competencies": ["BPMN"],
    "jobs": [
        {"title": "BA", "company": "Bank plc", "dates": "2020 – Present", "bullets": ["Did X."]}
    ],
}
TAILORED = {
    "headline": "Lead BA | Claims",
    "summary": "Business analyst with claims experience.",
    "competencies": ["BPMN", "Stakeholder management"],
    "jobs": PROFILE["jobs"],
    "achievements": ["Saved £1m."],
    "certifications": "",
    "education": "BSc",
    "keywords_used": [],
    "changes": [],
    "gaps": ["Guidewire"],
    "honesty_check": "Nothing invented.",
}


class FakeServer:
    def __init__(self):
        self.jobs, self.posts, self.results = {}, [], {}
        fake = self

        class Handler(BaseHTTPRequestHandler):
            def log_message(self, *args):
                pass

            def _send(self, code, body):
                raw = json.dumps(body).encode()
                self.send_response(code)
                self.send_header("Content-Type", "application/json")
                self.end_headers()
                self.wfile.write(raw)

            def do_GET(self):
                if self.headers.get("Authorization") != "Bearer cwk_test":
                    return self._send(401, {"detail": "invalid or revoked API key"})
                if self.path == "/api/v1/profiles":
                    return self._send(
                        200, {"profiles": [{"name": "cv_tailor"}, {"name": "cv_extract"}]}
                    )
                job = fake.jobs.get(self.path.rsplit("/", 1)[-1])
                return (
                    self._send(200, {"job": job})
                    if job
                    else self._send(404, {"detail": "job not found"})
                )

            def do_POST(self):
                body = json.loads(self.rfile.read(int(self.headers["Content-Length"])) or b"{}")
                if self.path.endswith("/cancel"):
                    return self._send(200, {"job": {}})
                if body["profile"] == "forbidden":
                    return self._send(403, {"detail": "this key may not use profile 'forbidden'"})
                fake.posts.append(body)
                job_id = f"job_{len(fake.posts)}"
                fake.jobs[job_id] = {
                    "id": job_id,
                    "status": "queued",
                    "progress": {"last": "Queued"},
                }
                return self._send(202, {"created": True, "job": fake.jobs[job_id]})

        self.httpd = HTTPServer(("127.0.0.1", 0), Handler)
        threading.Thread(target=self.httpd.serve_forever, daemon=True).start()
        self.url = f"http://127.0.0.1:{self.httpd.server_port}"

    def finish(self, job_id, result, status="succeeded", error=None):
        self.jobs[job_id].update(
            status=status,
            result=result,
            error=error,
            usage={"cost_usd_estimate": 0.02},
            progress={"last": "Finished"},
        )


@pytest.fixture
def env(tmp_path, monkeypatch):
    server = FakeServer()
    monkeypatch.setattr(config, "CLAUDE_WEB_URL", server.url)
    monkeypatch.setattr(config, "CLAUDE_WEB_API_KEY", "cwk_test")
    monkeypatch.setattr(config, "PROFILE_PATH", tmp_path / "profile.json")
    monkeypatch.setattr(config, "CV_DIR", tmp_path / "cvs")
    monkeypatch.setattr(config, "COVER_LETTER_DIR", tmp_path / "letters")
    monkeypatch.setattr(config, "UPLOAD_DIR", tmp_path / "uploads")
    path = tmp_path / "t.db"
    db.init_db(path)
    with db.connect(path) as conn:
        yield {"server": server, "conn": conn, "tmp": tmp_path}
    server.httpd.shutdown()


def _role(conn, **extra):
    data = {
        "title": "Lead BA",
        "company": "Insurer",
        "link": "https://jobs.example/99",
        "status": "Found",
        **extra,
    }
    rid = db.insert_job(conn, data)
    conn.commit()
    return rid


def test_disabled_without_key(monkeypatch):
    monkeypatch.setattr(config, "CLAUDE_WEB_API_KEY", "")
    assert not cj.enabled()
    with pytest.raises(cj.ClaudeWebError, match="isn't connected"):
        cj.server_profiles()


def test_errors_are_readable(env, monkeypatch):
    assert cj.server_profiles() == ["cv_tailor", "cv_extract"]
    with pytest.raises(cj.ClaudeWebError, match="403: this key may not"):
        cj.submit(env["conn"], "forbidden", {})
    monkeypatch.setattr(config, "CLAUDE_WEB_API_KEY", "cwk_wrong")
    with pytest.raises(cj.ClaudeWebError, match="401") as err:
        cj.server_profiles()
    assert err.value.status == 401


def test_candidate_prefers_profile_then_uploaded_cv(env):
    assert cj.load_candidate() == (None, "")
    cv_upload.save_upload("cv.txt", ("Ada Example, business analyst. " * 20).encode())
    cand, source = cj.load_candidate()
    assert source == "cv" and cand["cv_text"].startswith("Ada Example")
    cj.save_profile({**PROFILE, "_comment": "dropped"})
    cand, source = cj.load_candidate()
    assert source == "profile" and cand == PROFILE


def test_role_input_cleans_nan_and_adds_advert():
    role = cj.role_input(
        {"title": "BA", "company": float("nan"), "link": None, "rate": 500}, "  JD text "
    )
    assert role["company"] == "" and role["link"] == "" and role["rate"] == "500"
    assert role["description"] == "JD text"
    assert cj.role_input({"title": ""})["title"] == "Untitled role"


def test_tailored_cv_is_filed_and_settles_a_queued_role(env):
    conn, server = env["conn"], env["server"]
    cj.save_profile(PROFILE)
    rid = _role(conn, status="Draft CV")
    ai_id = cj.submit(conn, "cv_tailor", {"candidate": PROFILE, "role": {"title": "Lead BA"}}, rid)
    assert server.posts[0]["profile"] == "cv_tailor"
    assert server.posts[0]["idempotency_key"].startswith("jobtracker-")
    assert cj.has_active(conn)

    assert cj.sync(conn) == []  # still queued
    server.finish("job_1", TAILORED)
    messages = cj.sync(conn)
    assert messages == ["Tailored CV: saved Ada Example - Lead BA (Claude).docx"]
    row = dict(conn.execute("SELECT * FROM jobs WHERE id = ?", (rid,)).fetchone())
    assert row["status"] == "CV Drafted" and row["cv_version"].endswith("(Claude).docx")
    doc = Document(str(env["tmp"] / "cvs" / row["cv_version"]))
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Ada Example" and texts[1] == "Lead BA | Claims"
    assert "London  |  ada@example.com" in texts  # contact comes from the profile
    assert not cj.has_active(conn)
    # A second sync / apply does nothing more.
    assert cj.sync(conn) == []
    assert cj.apply(conn, ai_id) == "saved Ada Example - Lead BA (Claude).docx"
    assert len(list((env["tmp"] / "cvs").iterdir())) == 1


def test_cover_letter_is_filed_without_touching_status(env):
    conn, server = env["conn"], env["server"]
    rid = _role(conn, status="Applied")
    cj.submit(conn, "cover_letter", {"candidate": PROFILE, "role": {"title": "Lead BA"}}, rid)
    server.finish(
        "job_1",
        {
            "subject": "Lead BA",
            "greeting": "Dear Sam,",
            "paragraphs": ["One.", "Two."],
            "sign_off": "Kind regards,\nAda",
        },
    )
    cj.sync(conn)
    row = dict(conn.execute("SELECT * FROM jobs WHERE id = ?", (rid,)).fetchone())
    assert row["status"] == "Applied"
    assert row["cover_letter"] == "Cover Letter - Insurer - Lead BA (Claude).docx"
    texts = [p.text for p in Document(str(env["tmp"] / "letters" / row["cover_letter"])).paragraphs]
    assert "Dear Sam," in texts and "Two." in texts and texts[-1] == "Ada"


def test_fit_review_waits_for_a_click(env):
    conn, server = env["conn"], env["server"]
    rid = _role(conn, fit_notes="old")
    ai_id = cj.submit(conn, "job_fit_review", {}, rid)
    server.finish(
        "job_1", {"fit_score": 72, "recommendation": "apply", "fit_notes": "Strong claims match"}
    )
    assert cj.sync(conn) == ["Fit review: succeeded"]
    assert conn.execute("SELECT fit_notes FROM jobs WHERE id = ?", (rid,)).fetchone()[0] == "old"
    assert cj.apply(conn, ai_id) == "fit notes updated"
    notes = conn.execute("SELECT fit_notes FROM jobs WHERE id = ?", (rid,)).fetchone()[0]
    assert notes == "72/100 apply: Strong claims match"


def test_deleted_role_and_failed_job(env):
    conn, server = env["conn"], env["server"]
    rid = _role(conn)
    cj.submit(conn, "cv_tailor", {}, rid)
    cj.submit(conn, "cover_letter", {}, rid)
    db.delete_job(conn, rid)
    conn.commit()
    server.finish("job_1", TAILORED)
    server.finish("job_2", None, status="failed", error="timed out after 10 minutes")
    messages = cj.sync(conn)
    assert "Tailored CV: the role was deleted, so nothing was filed" in messages
    assert "Cover letter: failed (timed out after 10 minutes)" in messages
    with pytest.raises(ValueError):
        cj.apply(conn, 2)


def test_job_gone_from_server_is_marked_failed(env):
    conn = env["conn"]
    cj.submit(conn, "cv_tailor", {}, None)
    env["server"].jobs.clear()
    assert cj.sync(conn) == ["Tailored CV: failed (the server no longer has this job)"]


def test_search_adds_only_chosen_untracked_roles(env):
    conn, server = env["conn"], env["server"]
    _role(conn, link="https://jobs.example/1", source_job_id="1")
    links, ids = cj.tracked_keys(conn)
    assert links == ["https://jobs.example/1"] and ids == ["1"]
    ai_id = cj.submit(conn, "job_search", {}, None)
    roles = [
        {"title": "Old", "link": "https://jobs.example/1", "source_job_id": "1", "fit_score": 90},
        {
            "title": "New A",
            "link": "https://jobs.example/2",
            "source_job_id": "2",
            "fit_score": 80,
            "fit_notes": "Good",
            "company": "Acme",
            "rate": "£550",
        },
        {"title": "New B", "link": "https://jobs.example/3", "source_job_id": "3", "fit_score": 60},
        {
            "title": "Same id",
            "link": "https://other.example/2",
            "source_job_id": "2",
            "fit_score": 70,
        },
    ]
    server.finish("job_1", {"summary": "", "roles": roles, "skipped": []})
    cj.sync(conn)
    chosen = ["https://jobs.example/1", "https://jobs.example/2", "https://other.example/2"]
    assert cj.apply(conn, ai_id, links=chosen) == "added 1 role(s)"
    added = dict(
        conn.execute("SELECT * FROM jobs WHERE link = ?", ("https://jobs.example/2",)).fetchone()
    )
    assert (
        added["status"] == "Found"
        and added["fit_notes"] == "80/100: Good"
        and added["rate"] == "£550"
    )
    # Adding more later from the same search works; repeats are skipped.
    assert (
        cj.apply(conn, ai_id, links=["https://jobs.example/2", "https://jobs.example/3"])
        == "added 1 role(s)"
    )
    assert conn.execute("SELECT COUNT(*) FROM jobs").fetchone()[0] == 3


def test_profile_from_cv_saves_with_backup(env):
    conn, server = env["conn"], env["server"]
    cj.save_profile({"name": "Old"})
    ai_id = cj.submit(conn, "cv_extract", {"cv_text": "x"}, None)
    server.finish("job_1", {**PROFILE, "warnings": ["Education has no dates."]})
    assert cj.sync(conn) == ["Profile from CV: succeeded"]  # never replaces a profile unasked
    assert json.loads(config.PROFILE_PATH.read_text())["name"] == "Old"
    assert cj.apply(conn, ai_id) == "saved profile.json"
    assert json.loads(config.PROFILE_PATH.read_text()) == PROFILE
    backups = list(env["tmp"].glob("profile.*.bak.json"))
    assert len(backups) == 1 and json.loads(backups[0].read_text()) == {"name": "Old"}


def test_cancel(env):
    conn = env["conn"]
    ai_id = cj.submit(conn, "job_search", {}, None)
    cj.cancel(conn, ai_id)
    assert cj.get(conn, ai_id)["status"] == "cancelled"
