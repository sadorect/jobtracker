"""Upload your existing CV and pull its text out.

The original file is kept in ``config.UPLOAD_DIR`` (timestamped, never
overwritten) and the most recent upload becomes the *current CV*. Its text is
what Claude turns into ``profile.json`` (claude-web ``cv_extract``), and what the
AI tools fall back to when there is no profile yet.

Supported: .docx, .pdf (text-based, not scanned images), .txt and .md.
"""

from __future__ import annotations

import datetime as dt
import io
import json
import re
from pathlib import Path

from docx import Document

import config

ALLOWED_SUFFIXES = (".docx", ".pdf", ".txt", ".md")
MAX_BYTES = 5 * 1024 * 1024
MIN_TEXT_CHARS = 200  # less than this is almost certainly a scan or an empty file
_CURRENT = "current_cv.json"


class CvUploadError(ValueError):
    """The file can't be used as a CV (wrong type, too big, no readable text)."""


def _docx_text(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # many CV templates put skills or dates in tables
        for row in table.rows:
            cells = []
            for cell in row.cells:
                text = cell.text.strip()
                if text and text not in cells:  # merged cells repeat their text
                    cells.append(text)
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _pdf_text(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - pypdf is in requirements.txt
        raise CvUploadError("PDF support needs the pypdf package (pip install pypdf).") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as exc:  # noqa: BLE001 - pypdf raises many types for bad files
        raise CvUploadError(f"Couldn't read that PDF: {exc}") from exc


def extract_text(filename: str, data: bytes) -> str:
    """Plain text of a CV file. Raises CvUploadError if it can't be used."""
    suffix = Path(filename or "").suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise CvUploadError(
            f"Upload a {', '.join(ALLOWED_SUFFIXES)} file (got '{suffix or 'no extension'}')."
        )
    if len(data) > MAX_BYTES:
        raise CvUploadError(f"That file is over {MAX_BYTES // (1024 * 1024)} MB.")
    if suffix == ".docx":
        try:
            text = _docx_text(data)
        except Exception as exc:  # noqa: BLE001 - not a real .docx
            raise CvUploadError(f"Couldn't read that Word file: {exc}") from exc
    elif suffix == ".pdf":
        text = _pdf_text(data)
    else:
        text = data.decode("utf-8", errors="replace")
    text = re.sub(r"[ \t ]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text).strip()
    if len(text) < MIN_TEXT_CHARS:
        raise CvUploadError(
            "Hardly any text came out of that file. If it's a scanned PDF, upload the "
            "Word version or a text-based PDF instead."
        )
    return text


def save_upload(filename: str, data: bytes, upload_dir: Path | None = None) -> dict:
    """Store the upload, make it the current CV, and return its metadata."""
    text = extract_text(filename, data)
    upload_dir = Path(upload_dir or config.UPLOAD_DIR)
    upload_dir.mkdir(parents=True, exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y%m%d-%H%M%S")
    suffix = Path(filename).suffix.lower()
    stored = upload_dir / f"cv-{stamp}{suffix}"
    text_path = stored.with_suffix(".extracted.txt")
    stored.write_bytes(data)
    text_path.write_text(text, encoding="utf-8")
    meta = {
        "file": stored.name,
        "text_file": text_path.name,
        "original_name": Path(filename).name,
        "uploaded_at": dt.datetime.now().isoformat(timespec="seconds"),
        "chars": len(text),
    }
    (upload_dir / _CURRENT).write_text(json.dumps(meta, indent=1), encoding="utf-8")
    return {**meta, "text": text}


def current_cv(upload_dir: Path | None = None) -> dict | None:
    """The most recent upload (metadata + ``text``), or None."""
    upload_dir = Path(upload_dir or config.UPLOAD_DIR)
    try:
        meta = json.loads((upload_dir / _CURRENT).read_text(encoding="utf-8"))
        meta["text"] = (upload_dir / meta["text_file"]).read_text(encoding="utf-8")
        meta["path"] = upload_dir / meta["file"]
        return meta
    except (OSError, ValueError, KeyError):
        return None
